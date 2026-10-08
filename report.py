import os
from datetime import datetime

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from config import DATA_DIR, REPORT_DIR


def load_data():
    latest = pd.read_csv(os.path.join(DATA_DIR, "books_latest.csv"))
    history = pd.read_csv(os.path.join(DATA_DIR, "books_history.csv"))
    return latest, history


def build_summary(latest, history):
    cheapest = latest.loc[latest["price"].idxmin()]
    priciest = latest.loc[latest["price"].idxmax()]

    return pd.DataFrame({
        "Metric": [
            "Report generated",
            "Books in latest scrape",
            "Scrape runs so far",
            "Average price (£)",
            "Cheapest book",
            "Cheapest price (£)",
            "Most expensive book",
            "Most expensive price (£)",
            "Average rating (out of 5)",
            "Books in stock",
        ],
        "Value": [
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            len(latest),
            int(history["scraped_at"].nunique()),
            round(float(latest["price"].mean()), 2),
            cheapest["title"],
            float(cheapest["price"]),
            priciest["title"],
            float(priciest["price"]),
            round(float(latest["rating"].mean()), 2),
            int(latest["in_stock"].sum()),
        ],
    })


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"

    for col in ws.columns:
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(longest + 3, 60)


def build_report(latest, history):
    os.makedirs(REPORT_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    path = os.path.join(REPORT_DIR, f"books_report_{stamp}.xlsx")

    summary = build_summary(latest, history)
    books = latest[["title", "price", "rating", "in_stock", "scraped_at"]].sort_values("price")
    ratings = latest.groupby("rating").size().reset_index(name="books")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        books.to_excel(writer, sheet_name="Books", index=False)
        ratings.to_excel(writer, sheet_name="Ratings", index=False)

        for ws in writer.book.worksheets:
            style_sheet(ws)

        # Bar chart: number of books per star rating
        ws = writer.sheets["Ratings"]
        chart = BarChart()
        chart.title = "Books by star rating"
        chart.x_axis.title = "Stars"
        chart.y_axis.title = "Number of books"
        data = Reference(ws, min_col=2, min_row=1, max_row=len(ratings) + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=len(ratings) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "D2")

    return path

